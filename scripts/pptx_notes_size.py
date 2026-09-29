"""Make the speaker notes of a rendered deck readable: an explicit font size on every
notes run (and paragraph end), shrink-to-fit switched off on the notes box, and the
notes master's default raised to the same size, so notes typed later match.

    python scripts/pptx_notes_size.py docs/slides/MNF15-talk-2026-09-v4.pptx [18]

pandoc writes notes with no size of their own, so they fall back to the reference
template's 12 pt notes style, which reads as tiny in PowerPoint's notes pane.
"""
import sys

from lxml import etree
from pptx import Presentation

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
AUTOFIT = (A + "normAutofit", A + "spAutoFit", A + "noAutofit")


def size_paragraph(p, sz):
    for r in p.findall(A + "r"):
        rpr = r.find(A + "rPr")
        if rpr is None:
            rpr = etree.SubElement(r, A + "rPr")
            r.remove(rpr)
            r.insert(0, rpr)
        rpr.set("sz", str(sz))
    end = p.find(A + "endParaRPr")
    if end is None:
        end = etree.SubElement(p, A + "endParaRPr")
    end.set("sz", str(sz))


def main(path, pt=18):
    sz = int(round(pt * 100))
    prs = Presentation(path)
    n = 0
    for s in prs.slides:
        if not s.has_notes_slide:
            continue
        tf = s.notes_slide.notes_text_frame
        if tf is None:
            continue
        body = tf._txBody
        bp = body.find(A + "bodyPr")
        for c in list(bp):
            if c.tag in AUTOFIT:
                bp.remove(c)
        etree.SubElement(bp, A + "noAutofit")
        for p in body.findall(A + "p"):
            size_paragraph(p, sz)
        n += 1
    style = prs.notes_master._element.find(P + "notesStyle")
    if style is not None:
        for d in style.iter(A + "defRPr"):
            d.set("sz", str(sz))
    prs.save(path)
    print(f"{path}: notes on {n} slides set to {pt:g} pt, shrink-to-fit off, notes master default {pt:g} pt")


if __name__ == "__main__":
    main(sys.argv[1], float(sys.argv[2]) if len(sys.argv) > 2 else 18)
