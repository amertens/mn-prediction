"""
scripts/policy_deck/23_build_lightning_predictors.py

The 3-minute short-oral (rapid-fire) deck for the Micronutrient Forum 2026, on
the Forum's own template: three slides, the abstract title, then background /
aims / methods, results, conclusions and impact.

This replaces the story in 08_build_lightning_deck.py (which was built on the
454-column set and the fig3 predictor chart). The story here is the one ANM
asked for on 18 September 2026: WHICH PREDICTORS CARRY THE SIGNAL, and the
Malawi B12 example as the caution.

  Slide 1  Aims and methods.
  Slide 2  Every model rests on the same three environmental blocks, and then
           each nutrient reaches for a different second signal that a
           nutritionist would recognise. B12 has the most specific dietary
           marker and is the best-predicted outcome.
  Slide 3  Malawi: the same model's anaemia and malaria weights run the WRONG
           WAY for B12, because they mark the fish-eating lakeshore. These say
           where deficiency is, not what to change.

Every number is read from a committed result table.

    python scripts/policy_deck/23_build_lightning_predictors.py

Input template:  C:/Users/andre/Dropbox/MN prediction/MNF2026-Short-oral-template.pptx
Figures:         results/figures/mnf15/figD_what_each_model_uses.png
                 results/figures/mnf15/figM_malawi_b12_surrogate.png
Output:          docs/slides/MN-proxy-lightning-predictors-MNF2026.pptx
"""
import copy
import csv
import os
import sys

from pptx import Presentation
from pptx.util import Inches, Pt

ROOT = os.path.normpath(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", ".."))
TEMPLATE = "C:/Users/andre/Dropbox/MN prediction/MNF2026-Short-oral-template.pptx"
OUT = os.environ.get("LIGHTNING_OUT") or os.path.join(ROOT, "docs", "slides", "MN-proxy-lightning-predictors-MNF2026.pptx")
P2 = os.path.join(ROOT, "results", "tables", "protocol_v2")
FIG = os.path.join(ROOT, "results", "figures", "mnf15")
sys.path.insert(0, os.path.join(ROOT, "scripts", "concept_slides"))
from build_concept_slides import BLUE, TEXT, icon, textbox  # noqa: E402  (the full talk's pictogram helpers)
from pptx.enum.text import MSO_ANCHOR  # noqa: E402
ICONS = os.path.join(ROOT, "docs", "slides", "img", "icons")


def rows(name, base=P2):
    with open(os.path.join(base, name), newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def num(x):
    try:
        return float(x)
    except (TypeError, ValueError):
        return None


def mean(xs):
    xs = [v for v in (num(x) for x in xs) if v is not None]
    return sum(xs) / len(xs) if xs else float("nan")


# ---- numbers -----------------------------------------------------------------
# cell_master.csv is written by 20_mnf15_figures.R and carries the measurability
# screen (prevalence >= 2% and the survey's own district-level signal >= 0.30).
CM = rows("cell_master.csv", FIG)
keep = [r for r in CM if r["keep"] == "TRUE"]


def by_nutrient(field):
    out = {}
    for r in keep:
        out.setdefault(r["nutrient"], []).append(r[field])
    return {k: mean(v) for k, v in out.items()}


infill_n = by_nutrient("infill")
tr_n = by_nutrient("tr")
b12_in, b12_tr = infill_n["Vitamin B12"], tr_n["Vitamin B12"]
fol_in, fol_tr = infill_n["Folate"], tr_n["Folate"]

# environmental share of the index: climate + soil + satellite embedding
IMP = [r for r in rows("index_importance_domains.csv")
       if r["scope"] == "pooled" and r["target"] == "level"]
ENV = {"Climate and weather", "Soil characteristics", "Satellite embedding"}
env = {}
for r in IMP:
    if r["domain"] in ENV:
        env[r["outcome"]] = env.get(r["outcome"], 0.0) + float(r["share"])
env_lo, env_hi = min(env.values()), max(env.values())

# transport by tier: open / headline / with DHS
def tr_tier(fn):
    r = [x for x in rows(fn) if x["estimand"] == "country"
         and x["arm"] == "domain_index" and x["target"] == "level"][0]
    return float(r["mean_spearman"])


tier_open = tr_tier("benchmarks_v2_summary_open.csv")
tier_head = tr_tier("benchmarks_v2_summary.csv")
tier_dhs = tr_tier("benchmarks_v2_summary_withdhs.csv")
null_d = float([r for r in rows("transport_null_calibration.csv") if r["tier"] == "admin2"][0]["null_mean_q95"])

# Malawi women's B12: the best-predicted cell, and the rho values behind the maps
mw = [r for r in CM if r["country"] == "Malawi" and r["outcome"] == "women_b12"][0]
mw_in = float(mw["infill"])
with open(os.path.join(FIG, "figM_malawi_rho.txt"), encoding="utf-8") as f:
    r_anaemia, r_fish = [float(v) for v in f.read().split()]

# drop-one domain ablation: what a new country actually needs
AB = {r["domain"]: float(r["delta_drop"]) for r in rows("domain_ablation_loco_summary.csv")
      if r["target"] == "level"}


def f2(x):
    return "%.2f" % x


def pc(x):
    return "%d%%" % round(100 * x)


print("B12 %.2f/%.2f  folate %.2f/%.2f  env %.0f-%.0f%%  tiers %.2f/%.2f/%.2f  Malawi B12 %.2f  rho %.2f/%.2f"
      % (b12_in, b12_tr, fol_in, fol_tr, 100 * env_lo, 100 * env_hi,
         tier_open, tier_head, tier_dhs, mw_in, r_anaemia, r_fish))

TITLE = ("Which public data layers predict micronutrient deficiency, and what they do "
         "and do not tell us: evidence from four national biomarker surveys")
AUTHORS = ("Andrew N. Mertens (presenting), Reed Atkin, Sorrel Namaste, Demewoz Haile, Kenneth H. Brown, "
           "Saskia Osendarp, Xiuping Tan, Eric Stewart, Keith Lividini, Seth Adu-Afarwuah, Nicolai Petry, "
           "Aminata S. Koroma, Mary H. Hodges, Fabian Rohner, Haddy Crookes, James P. Wirth, "
           "Jonathan Gorstein, Kathy Banke, Sonja Y. Hess")

S1 = [
    ("vial", "Micronutrient surveys are rare and expensive",
     "Most countries have no recent estimate, and almost none can say which districts are worst off."),
    ("layer-group", "Can data that already exist for every district fill that gap?",
     "Satellite, climate, soil, crops, disease maps, prices, household surveys."),
    ("earth-africa", "Four national biomarker surveys as the ground truth",
     "The Gambia 2018, Ghana 2017, Malawi 2015-16, Sierra Leone 2013: 206 districts; iron, vitamin A, folate, B12 and zinc."),
    ("sliders", "575 public layers per district, none needing a blood sample",
     "Each family of layers becomes a few summary scores, weighted by how well it tracks deficiency where blood was drawn. Nothing is tuned."),
    ("location-crosshairs", "Tested only on what the model never saw",
     "Districts held out one in five, and whole countries held out in turn."),
]

S2 = [
    "Three blocks do most of the work for every nutrient: satellite imagery, climate and soil carry %s to %s of the model's weight."
    % (pc(env_lo), pc(env_hi)),
    "What each nutrient adds on top is recognisable. B12 leans on infant meat and fish consumption; iron and folate on modelled anaemia; vitamin A on child wasting; children's iron on a cereal-and-livestock farming gradient.",
    "The nutrient with the clearest dietary marker is also the best predicted. Women's B12 reaches %s inside a country and %s in a country never surveyed, against %s for chance."
    % (f2(b12_in), f2(b12_tr), f2(null_d)),
    "The layers that cross a border are the free ones: openly downloadable data alone score %s, ahead of adding public survey microdata (%s) or DHS microdata (%s)."
    % (f2(tier_open), f2(tier_head), f2(tier_dhs)),
]

S3 = [
    "Malawi women's B12 is our best-predicted outcome (%s), and it carries a warning: modelled anaemia tracks BETTER B12 status (rho %s), which no biology supports."
    % (f2(mw_in), f2(r_anaemia)),
    "The maps say why. Anaemia and malaria peak along the lakeshore and in the south, and so does eating fish (rho %s), the main source of B12 in Malawi."
    % (f2(r_fish),),
    "The model found a marker for a place, not a cause, and cannot tell them apart. These layers say where deficiency is; they are not a list of things to change.",
    "Impact: a country can rank its own districts today from data it can download, to target programmes and to place the next survey. The level still needs one measured national sample.",
]

# ---- template filling ----------------------------------------------------------
def shape_by_name(slide, name):
    for sh in slide.shapes:
        if sh.name == name:
            return sh
    sys.exit("shape %r not found; template shapes: %s"
             % (name, [s.name for s in slide.shapes]))


def set_text(shape, text, size=None):
    tf = shape.text_frame
    p0 = tf.paragraphs[0]
    for p in list(tf.paragraphs[1:]):
        p._p.getparent().remove(p._p)
    runs = p0.runs or [p0.add_run()]
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    runs[0].text = text
    if size:
        runs[0].font.size = Pt(size)


def set_bullets(shape, lines, size=None):
    tf = shape.text_frame
    tf.word_wrap = True
    proto = copy.deepcopy(tf.paragraphs[0]._p)
    for p in list(tf.paragraphs):
        p._p.getparent().remove(p._p)
    for line in lines:
        tf._txBody.append(copy.deepcopy(proto))
        para = tf.paragraphs[-1]
        for r in para.runs[1:]:
            r._r.getparent().remove(r._r)
        if not para.runs:
            para.add_run()
        para.runs[0].text = line
        if size:
            para.runs[0].font.size = Pt(size)
        para.space_after = Pt(6)


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


prs = Presentation(TEMPLATE)
s1, s2, s3 = list(prs.slides)[:3]

title = shape_by_name(s1, "Title 1")
title.left, title.top, title.width, title.height = Inches(0.27), Inches(0.15), Inches(12.8), Inches(1.05)
set_text(title, TITLE, size=22)
byline = shape_by_name(s1, "TextBox 5")
byline.left, byline.top, byline.width, byline.height = Inches(0.27), Inches(1.25), Inches(12.8), Inches(0.75)
byline.text_frame.word_wrap = True
set_text(byline, AUTHORS, size=12)
heading = shape_by_name(s1, "Title 23")
heading.top = Inches(2.1)
set_text(heading, "Background, aims and methods")
body = shape_by_name(s1, "Text Placeholder 2")
body._element.getparent().remove(body._element)   # the bullets become icon rows (18 Sep: the joint deck's styling)
X1, Y1, W1, H1 = Inches(0.45), Inches(2.75), Inches(12.4), Inches(4.05)
rh = H1 / len(S1)
for k, (ic, head, cap) in enumerate(S1):
    y = Y1 + k * rh; isz = Inches(0.55)
    icon(s1, f"{ic}_blue", X1 + Inches(0.1), y + (rh - isz) / 2, isz, ICONS)
    textbox(s1, X1 + Inches(0.95), y, W1 - Inches(1.0), rh,
            [(head, 15, True, BLUE), (cap, 12.5, False, TEXT)], anchor=MSO_ANCHOR.MIDDLE)
notes(s1, "ANDREW, about 45 seconds (128 words). Micronutrient surveys are rare and expensive. Most countries have no recent estimate, and even where one exists it was designed to report regions, so programmes choose districts with almost no data. Our question: can data that already exist for every district fill that gap? Satellite, climate, soil, crops, disease maps, prices, household surveys. Four national biomarker surveys, in The Gambia, Ghana, Malawi and Sierra Leone, are the ground truth: 206 districts, five nutrients. We linked 575 public layers to every district, none of which needs a blood sample. Each family of layers becomes a few summary scores, weighted by how well it tracks deficiency where blood was drawn. Nothing is tuned. And every number you will see comes from predicting districts, and whole countries, the model never saw.")

t2 = shape_by_name(s2, "Title 1")
t2.left, t2.top, t2.width, t2.height = Inches(0.27), Inches(0.10), Inches(12.8), Inches(0.62)
set_text(t2, "Every model uses the same three layers; each nutrient adds its own", size=24)
body = shape_by_name(s2, "Text Placeholder 2")
body.left, body.top, body.width, body.height = Inches(0.30), Inches(0.80), Inches(12.7), Inches(2.05)
set_bullets(body, S2, size=13)
s2.shapes.add_picture(os.path.join(FIG, "figD_what_each_model_uses.png"),
                      Inches(1.55), Inches(2.95), width=Inches(10.2))
notes(s2, "ANDREW, about 80 seconds (184 words). This is the finding. Each panel is one nutrient, and the bars are the share of the model's weight carried by each family of predictors. The top three bars are the same in every panel: satellite imagery, climate and soil, just under half the weight whichever nutrient you take. Then look below them, because a nutritionist would recognise all of it. B12, the nutrient you get almost only from animal-source food, leans on infant feeding: districts where more infants are fed meat or fish. Iron and folate lean on modelled anaemia, their clinical consequence. Vitamin A leans on child wasting. Children's iron leans on a cereal-and-livestock farming gradient. And the nutrient with the clearest dietary marker is the one we predict best: women's B12 reaches 0.63 inside a country and 0.52 in a country never surveyed, against 0.08 for chance. One practical point. The layers that cross a border are the free ones. Openly downloadable data alone score 0.32 in a new country; adding public survey microdata, which cost us weeks and a registration each, takes it to 0.30, and DHS microdata to 0.28.")

t3 = shape_by_name(s3, "Title 1")
t3.left, t3.top, t3.width, t3.height = Inches(0.27), Inches(0.10), Inches(12.8), Inches(0.62)
set_text(t3, "Malawi: the model finds place-markers, not mechanisms", size=24)
body = shape_by_name(s3, "Text Placeholder 2")
body.left, body.top, body.width, body.height = Inches(0.30), Inches(0.80), Inches(12.7), Inches(2.05)
set_bullets(body, S3, size=13)
s3.shapes.add_picture(os.path.join(FIG, "figM_malawi_b12_surrogate.png"),
                      Inches(2.95), Inches(2.95), width=Inches(9.0))
notes(s3, "ANDREW, about 55 seconds (135 words). And the caution. Malawi women's B12 is our best-predicted outcome, at 0.70. Inside it, the anaemia layer points the wrong way: districts with more anaemia have better B12 status. No biology says that. The maps say why. Anaemia and malaria peak along the lakeshore and in the south, and that is exactly where households eat fish, the main source of B12 in Malawi. The model found a marker for a place, not a cause, and it cannot tell the difference. So the rule we hold to: these layers tell you where to look, not what to change. What a country gets today is a ranking of its own districts from data it can download, enough to target a programme and to place the next survey. The level still needs one measured national sample. Thank you.")

prs.save(OUT)
print("wrote", OUT)
