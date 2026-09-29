"""Copy into the rendered v5 MNF15 deck the slides that are not built from source: the
full-talk slides and Andrew's hand-edited slides, taken from his working copy
docs/slides/MNF15-talk-2026-09-v4-ANM.pptx (28 September; read only), each placed after
the v5 slide it follows in his order. The "full talk" tags are removed (his comment).

    python scripts/policy_deck/38_merge_v5_slides.py

Edits applied to copied slides (his comments on v4-ANM):
  15  prevalence error redrawn: an X for the national estimate, no blue rows, no text below
  18  Malawi four-panel: a caption that fits (appendix)
  22  tied to the previous slide: retitled, and the post-fix numbers of the per-nutrient chart
  24  post-fix model score, and the answer to "is this the average model / any simulation?"
The speaker notes are set to 18 pt afterwards (scripts/pptx_notes_size.py).
"""
import copy
import io
import os
import sys

from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = "C:/Users/andre/OneDrive/Documents/mn-prediction"
ANM = os.path.join(ROOT, "docs/slides/MNF15-talk-2026-09-v4-ANM.pptx")
DST = os.path.join(ROOT, "docs/slides/MNF15-talk-2026-09-v5.pptx")
FIG = os.path.join(ROOT, "results/figures/mnf15_v5")

# (v4-ANM slide number, v5 slide it follows; None = after the previous copy)
PLAN = [
    (6, "Which of these causes can public data measure?"),
    (12, "Why not a more complex machine-learning model?"), (13, None),
    (15, "But can proxy models recover national-level prevalence estimates?"),
    (22, "How well does it rank districts, nutrient by nutrient?"), (24, None),
    # appendix
    (9, "How does the model work, and how was it tested?"),
    (18, "How well does it rank districts inside a surveyed country?"),
    (21, "What happens when a whole region is held out?"),
    (26, "How do proxy models compare with the survey's own information?"),
    (37, "Is it just geography? Comparison with DHS-style geostatistics"), (38, None),
    (42, "Which data sources contribute most to model prediction performance?"),
    (44, "What can we learn about unused predictor domains?"), (45, None), (46, None), (47, None),
    (49, "Which outcomes and cut-offs does the analysis use?"), (50, None),
    (52, "Does targeting reach more deficient people?"),
    (54, "Accuracy pooled over all nutrients, in district pairs"),
    (59, "The live dashboard"),
]
PICTURE = {15: "v4_ft_prev_error.png", 18: "v4_ft_malawi_b12_four.png"}
TEXT = {  # paragraph text replacements (whole paragraph, first run keeps the formatting)
    22: {
        "What does a rank correlation of 0.40 mean on the ground?": "What does that accuracy mean on the ground?",
        "Take any two districts: the best model orders them the way the survey does 60% of the time":
            "Take any two districts: the model orders them the way the survey does 66% of the time",
        "Counted over every pair of districts. A random ordering gets 50% on average, the survey's own regional averages 56%; in a country the model has never seen, 55%.":
            "The average of the previous slide. A coin toss gets 50%, the survey's own regional averages 61%; in a country the model has never seen, 62%.",
        "In the fifth of districts it flags, deficiency runs at 32% against 24% nationally":
            "In the fifth of districts it flags, deficiency runs at 31% against 22% nationally",
        "A map that knew every district would find 41% in its worst fifth.": "A map that knew every district would find 47% in its worst fifth.",
        "That fifth holds 22% of the deficient people": "That fifth holds 25% of the deficient people",
        "The survey's own regional averages capture 20%; a perfect map 48%. Real, modest, and better than today's number.":
            "The survey's own regional averages capture 21%; a perfect map 40%. Real, modest, and better than today's number.",
    },
    24: {
        "The model reaches 59% of that ceiling": "The model reaches about 60% of that ceiling",
        "0.28 on the prevalence target against a ceiling of 0.48. The gap is mostly the survey's noise, not the model's ignorance.":
            "0.30 on the prevalence target against a ceiling of 0.48. The gap is mostly the survey's noise, not the model's ignorance.",
    },
}
NOTES = {  # replace the speaker notes of these copied slides
    15: ("ANDREW, about 40 seconds. Same combinations, now the level: how many percentage points each district's predicted prevalence "
         "is from the survey's, each district held out. At district level nobody has a good level: the survey's own regional averages "
         "are no better than one national number for every district. The model, shrunk for this job, is slightly better than both, "
         "and that is all. It is built to order districts, not to measure them.\n\n"
         "Backup, if asked: population-weighted mean absolute error of the held-out district prevalence (benchmarks_v2_cells.csv, "
         "prevalence target, post-fix). Means over 18 combinations: model (calibrated for level) 10.6 points, survey's regional average "
         "12.0, national estimate 11.7. Read the error against the outcome: 10 points on an outcome at 40% differs from 3 points on one "
         "at 4%, which is why rare outcomes look accurate here. The ranking uses the uncalibrated index; any level uses the calibrated one."),
    22: ("ANDREW, about 30 seconds. What does that accuracy mean on the ground? Take any two districts: the model orders them as the "
         "survey does 66 times in 100, the average of the previous slide; the survey's own regional averages 61, a coin 50. The fifth of "
         "districts it flags run at 31 per cent deficient against 22 nationally, and hold a quarter of the deficient people, against "
         "21 per cent for the regional averages' pick.\n\n"
         "Backup, if asked: post-fix, the 16 mappable nutrient-country combinations, each district held out (v3_percell_pairs.csv, "
         "nce_targeting_metrics.csv). Pairs inside a country: model 66 (14 combinations with an in-country test), regional average 61; "
         "held out 62 (16). Worst fifth: deficiency 30.9 against 21.7 per cent nationally; a perfect map finds 47.1; share of deficient "
         "people reached 25.1 (regional averages 21.0, perfect 40.3, a random fifth 20). Retitled on Andrew's comment ('why this number'): "
         "the earlier 0.40 was the pooled rank correlation, now tied to the pairs on the previous slide."),
}
NOTES_APPEND = {
    24: ("\n\nANSWER TO ANDREW'S COMMENT (28 Sep): not a simulation. The model score is the index (the deployed model), averaged over "
         "the 18 in-country nutrient-country combinations with each district held out (post-fix 0.30 on prevalence, 0.40 on average "
         "status). The ceiling is analytic: a variance-components model of respondents within clusters within districts, fitted to "
         "each survey (variance_components_ceiling.csv, 8 September build until the downstream rerun reaches it). The clusters-per-district "
         "projection (0.54, 0.66, 0.73) comes from the same variance model, not from simulated surveys. Post-fix: model 0.40 against the "
         "regional average 0.31, better in 13 of 18."),
}


def title_of(slide):
    t = slide.shapes.title
    return t.text_frame.text.strip().replace("\u2019", "'").replace("\u2018", "'") if t is not None else ""


def copy_slide(src_slide, dst):
    layout = next(l for l in dst.slide_layouts if l.name == src_slide.slide_layout.name)
    new = dst.slides.add_slide(layout)
    for ph in list(new.placeholders):
        ph._element.getparent().remove(ph._element)
    rid = {}
    for r_id, rel in src_slide.part.rels.items():
        if rel.reltype == RT.IMAGE:
            _, rid[r_id] = new.part.get_or_add_image_part(io.BytesIO(rel.target_part.blob))
        elif rel.reltype == RT.HYPERLINK and rel.is_external:
            rid[r_id] = new.part.relate_to(rel.target_ref, RT.HYPERLINK, is_external=True)
    tree = new.shapes._spTree
    for shp in src_slide.shapes:
        if shp.has_text_frame and shp.text_frame.text.strip() == "full talk":
            continue   # the tag goes (Andrew: "remove this full talk detail everywhere")
        el = copy.deepcopy(shp._element)
        for node in el.iter():
            for a in list(node.attrib):
                if a.split("}")[-1] in ("embed", "link", "id") and node.attrib[a] in rid:
                    node.attrib[a] = rid[node.attrib[a]]
        tree.append(el)
    if src_slide.has_notes_slide:
        new.notes_slide.notes_text_frame.text = src_slide.notes_slide.notes_text_frame.text
    return new


def move_after(dst, new_slide, after_idx):
    lst = dst.slides._sldIdLst
    el = next(e for e in lst if e.rId == dst.part.relate_to(new_slide.part, RT.SLIDE))
    lst.remove(el)
    lst.insert(after_idx + 1, el)


def index_of(dst, title):
    for i, s in enumerate(dst.slides):
        if title_of(s) == title:
            return i
    sys.exit(f"v5 slide not found: {title!r}")


def replace_picture(slide, png):
    from PIL import Image
    pic = max([s for s in slide.shapes if s.shape_type == 13], key=lambda p: p.width * p.height)
    left, top, w, h = pic.left, pic.top, pic.width, pic.height
    pic._element.getparent().remove(pic._element)
    iw, ih = Image.open(png).size
    sc = min(w / iw, h / ih)
    nw, nh = int(iw * sc), int(ih * sc)
    slide.shapes.add_picture(png, left + (w - nw) // 2, top + (h - nh) // 2, nw, nh)


def replace_text(slide, mapping):
    done = set()
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        for p in sh.text_frame.paragraphs:
            t = p.text.strip().replace("\u2019", "'")
            if t in mapping and p.runs:
                p.runs[0].text = mapping[t]
                for r in p.runs[1:]:
                    r.text = ""
                done.add(t)
    missing = set(mapping) - done
    if missing:
        sys.exit(f"text not found on slide: {sorted(missing)}")


def main():
    anm = list(Presentation(ANM).slides)
    dst = Presentation(DST)
    last = None
    for num, after in PLAN:
        new = copy_slide(anm[num - 1], dst)
        if num in PICTURE:
            replace_picture(new, os.path.join(FIG, PICTURE[num]))
        if num in TEXT:
            replace_text(new, TEXT[num])
        if num in NOTES:
            new.notes_slide.notes_text_frame.text = NOTES[num]
        if num in NOTES_APPEND:
            nt = new.notes_slide.notes_text_frame
            nt.text = nt.text + NOTES_APPEND[num]
        idx = last if after is None else index_of(dst, after)
        move_after(dst, new, idx)
        last = idx + 1
        print(f"  v4-ANM {num:2d} -> after {title_of(dst.slides[idx])[:60]!r}")
    dst.save(DST)
    print(f"v5: {len(dst.slides)} slides")
    sys.path.insert(0, os.path.join(ROOT, "scripts"))
    import pptx_notes_size
    pptx_notes_size.main(DST, 18)


if __name__ == "__main__":
    main()
