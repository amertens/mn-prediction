"""Set explicit font sizes on a rendered deck's body text and tables.

pandoc's PowerPoint writer leaves run sizes unset, so every bullet inherits the
reference document's master style (28 pt at the first level in the Ghana
template), which is too large for text-heavy slides. This post-processing step
stamps a size on every run outside the title placeholders: `--body` for bullets
and column text (one step smaller per indent level), `--table` for table cells.
Speaker notes are left alone.

    python scripts/pptx_set_text_size.py docs/slides/deck.pptx --body 18 --table 11

Called by scripts/render_deck.sh when it is given `--text BODY[,TABLE]`.
"""
import argparse
from pptx import Presentation
from pptx.util import Pt


def set_sizes(path, body, table):
    prs = Presentation(path)
    n_body = n_table = 0
    for slide in prs.slides:
        title = slide.shapes.title
        for shape in slide.shapes:
            if shape.has_text_frame and shape is not title and shape.text_frame.text.strip():
                if shape.is_placeholder and shape.placeholder_format.type in (1, 3):   # TITLE, CENTER_TITLE
                    continue
                for para in shape.text_frame.paragraphs:
                    size = max(body - 2 * (para.level or 0), 10)
                    for run in para.runs:
                        run.font.size = Pt(size); n_body += 1
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        for para in cell.text_frame.paragraphs:
                            for run in para.runs:
                                run.font.size = Pt(table); n_table += 1
    prs.save(path)
    print(f"{path}: {n_body} body run(s) set to {body} pt (minus 2 per indent level), {n_table} table run(s) set to {table} pt")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("pptx")
    ap.add_argument("--body", type=float, default=18)
    ap.add_argument("--table", type=float, default=11)
    a = ap.parse_args()
    set_sizes(a.pptx, a.body, a.table)
