"""Build the illustrated concept slides of a deck from a YAML spec, on the deck's
own reference template, as a standalone pptx the author pastes into the deck.

    python scripts/concept_slides/build_concept_slides.py docs/slides/MN-proxy-Ghana-concept-slides-2026-09.yaml
    python scripts/concept_slides/build_concept_slides.py docs/slides/deck.concept.yaml --into docs/slides/deck.pptx

Without --into, a standalone pptx of the slides is written next to the spec (with
previews and a placement table). With --into, the slides are drawn IN PLACE on a
deck pandoc has rendered: each spec entry is matched to the rendered slide with the
same title (the qmd keeps that slide as a title plus speaker notes), the empty body
placeholder is removed and the layout drawn; the title slide is matched by layout
and gets the logos and acknowledgement of a `title` entry. Called by
scripts/render_deck.sh when it is given `--concept <spec.yaml>`.

Reads, next to the spec: the deck qmd it names (for the reference document and
the speaker notes of the slides it replaces) and `<deck>.quantities.json`
(written by build_quantities.R) for every `{name}` placeholder. Writes
`<spec>.pptx`, one preview PNG per slide under `concept_previews/`, and
`<spec>-PLACEMENT.md`. Icons are read from docs/slides/img/icons/.

Layouts: pipeline (cards in a row with arrows), cards (cards with a heading and
lines), checklist (status tile, question, answer), icon_rows (icon, heading,
caption; optional image on the right), two_panel (two headed columns of icon
items), grid (groups of items in two columns).
"""
import copy
import json
import os
import re
import string
import subprocess
import sys

import yaml
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

BLUE = RGBColor(0x1F, 0x4E, 0x79)
LIGHT = RGBColor(0x6B, 0xAE, 0xD6)
CARD = RGBColor(0xEE, 0xF3, 0xF8)
TEXT = RGBColor(0x33, 0x33, 0x33)
MUTED = RGBColor(0x66, 0x66, 0x66)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
STATUS = {"yes": ("check_green", RGBColor(0xE6, 0xF2, 0xE8)), "partly": ("triangle-exclamation_amber", RGBColor(0xFC, 0xF3, 0xDC)),
          "no": ("xmark_grey", RGBColor(0xEC, 0xEC, 0xEC))}
# content area of the Ghana template: below the title, above the footer band
X0, Y0, W, H = Inches(0.92), Inches(1.85), Inches(11.5), Inches(4.6)


class QFormatter(string.Formatter):
    """{x:.2f} as usual; {x:pc} percent; {x:sgn} signed; missing -> 'NA'."""
    def get_value(self, key, args, kwargs):
        return kwargs.get(key, None) if isinstance(key, str) else super().get_value(key, args, kwargs)

    def format_field(self, value, spec):
        if value is None:
            return "NA"
        if spec == "pc":
            return f"{100 * value:.0f}%"
        if spec == "sgn":
            return f"{value:+.2f}"
        if spec == "" and isinstance(value, float):
            return f"{value:.2f}"
        return super().format_field(value, spec)


FMT = QFormatter()


def fill_text(s, Q):
    return FMT.vformat(str(s), (), Q) if s is not None else ""


# ---- notes: the replaced slide's speaker notes, inline R resolved from Q ------------------
def deck_notes(qmd_lines, title, Q):
    heads = [i for i, l in enumerate(qmd_lines) if l.strip() == f"## {title}"]
    if not heads:
        return ""
    i = heads[0] + 1
    while i < len(qmd_lines) and not qmd_lines[i].startswith("## ") and not qmd_lines[i].startswith("# "):
        if qmd_lines[i].strip() == "::: {.notes}":
            j = i + 1
            while qmd_lines[j].strip() != ":::":
                j += 1
            return resolve_inline_r(" ".join(l.strip() for l in qmd_lines[i + 1:j]), Q)
        i += 1
    return ""


def resolve_inline_r(text, Q):
    def q(name, idx=None):
        v = Q.get(name)
        if isinstance(v, dict) and idx:
            v = v.get(idx)
        return v

    def rep(m):
        expr = m.group(1).strip()
        mm = re.match(r'^(f2|pc|sgn)\(Q\$(\w+)(?:\["(\w)"\])?(?:,\s*(\d))?\)$', expr)
        if mm:
            fn, name, idx, d = mm.groups(); v = q(name, idx)
            if v is None:
                return "NA"
            d = int(d) if d else {"f2": 2, "pc": 0, "sgn": 2}[fn]
            return {"f2": f"{v:.{d}f}", "pc": f"{100 * v:.{d}f}%", "sgn": f"{v:+.{d}f}"}[fn]
        mm = re.match(r'^Q\$(\w+)(?:\["(\w)"\])?$', expr)
        if mm:
            v = q(*mm.groups())
            return "NA" if v is None else (f"{v:.0f}" if isinstance(v, float) and v == int(v) else str(v))
        return m.group(0)
    return re.sub(r"`r ([^`]+)`", rep, text)


# ---- drawing helpers -----------------------------------------------------------------------
def rounded_box(slide, x, y, w, h, fill=CARD, radius=0.12):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    sh.adjustments[0] = min(radius * Inches(1) / min(w, h), 0.5)
    sh.fill.solid(); sh.fill.fore_color.rgb = fill; sh.line.fill.background(); sh.shadow.inherit = False
    sh.text_frame.text = ""
    return sh


def textbox(slide, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT):
    """paras: list of (text, size_pt, bold, colour)."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05); tf.margin_top = tf.margin_bottom = Inches(0.03)
    for k, (text, size, bold, colour) in enumerate(paras):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        p.alignment = align
        if k:
            p.space_before = Pt(3)
        r = p.add_run(); r.text = text
        r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = colour; r.font.name = "Calibri"
    return tb


def icon(slide, name, x, y, size, icon_dir):
    f = os.path.join(icon_dir, name + ".png")
    if not os.path.exists(f):
        sys.exit(f"missing icon {f}: run build_quantities.R first")
    return slide.shapes.add_picture(f, x, y, height=size)


def icon_file(item, colour="blue"):
    nm = item.get("icon", "circle")
    return nm.replace("@", "_") if "@" in nm else f"{nm}_{colour}"


# ---- layouts -------------------------------------------------------------------------------
def layout_pipeline(slide, spec, Q, icons):
    items = spec["items"]; n = len(items)
    ncol = 3 if n > 4 else n; nrow = -(-n // ncol)
    gap = Inches(0.55); cw = (W - gap * (ncol - 1)) / ncol; ch = (H - Inches(0.3) * (nrow - 1)) / nrow
    if ncol >= 4 and nrow == 1:
        ch = min(ch, Inches(3.6))
    for k, it in enumerate(items):
        r, c = divmod(k, ncol)
        x = X0 + c * (cw + gap); y = Y0 + r * (ch + Inches(0.3))
        rounded_box(slide, x, y, cw, ch)
        if ncol >= 4:   # narrow cards: icon on top, heading and caption below it
            isz = Inches(0.6)
            icon(slide, icon_file(it), x + (cw - isz) / 2, y + Inches(0.2), isz, icons)
            textbox(slide, x + Inches(0.12), y + Inches(0.85), cw - Inches(0.24), Inches(0.8),
                    [(f"{k + 1}. {fill_text(it['heading'], Q)}", 15, True, BLUE)], anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
            textbox(slide, x + Inches(0.15), y + Inches(1.7), cw - Inches(0.3), ch - Inches(1.8),
                    [(fill_text(it.get("caption", ""), Q), 12.5, False, TEXT)])
        else:
            isz = Inches(0.62)
            icon(slide, icon_file(it), x + Inches(0.2), y + Inches(0.18), isz, icons)
            textbox(slide, x + Inches(0.95), y + Inches(0.12), cw - Inches(1.05), Inches(0.75),
                    [(f"{k + 1}. {fill_text(it['heading'], Q)}", 17, True, BLUE)], anchor=MSO_ANCHOR.MIDDLE)
            textbox(slide, x + Inches(0.2), y + Inches(0.92), cw - Inches(0.35), ch - Inches(1.0),
                    [(fill_text(it.get("caption", ""), Q), 13, False, TEXT)])
        if c < ncol - 1 and k < n - 1:
            ar = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, x + cw + Inches(0.14), y + ch / 2 - Inches(0.16), Inches(0.28), Inches(0.32))
            ar.fill.solid(); ar.fill.fore_color.rgb = LIGHT; ar.line.fill.background()


def layout_cards(slide, spec, Q, icons):
    items = spec["items"]; n = len(items)
    hs = spec.get("heading_size", 20); ts = spec.get("text_size", 14); isz = Inches(spec.get("icon_size", 1.0))
    gap = Inches(0.45 if n <= 3 else 0.25); cw = (W - gap * (n - 1)) / n; ch = H
    pad = Inches(0.3 if n <= 3 else 0.15)
    for k, it in enumerate(items):
        x = X0 + k * (cw + gap); y = Y0
        rounded_box(slide, x, y, cw, ch)
        icon(slide, icon_file(it), x + (cw - isz) / 2, y + Inches(0.25), isz, icons)
        hy = y + Inches(0.3) + isz + Inches(0.1)
        textbox(slide, x + Inches(0.1), hy, cw - Inches(0.2), Inches(0.5), [(fill_text(it["heading"], Q), hs, True, BLUE)], align=PP_ALIGN.CENTER)
        ty = hy + Inches(0.6)
        tb = textbox(slide, x + pad, ty, cw - 2 * pad, ch - (ty - y) - Inches(0.1), [(fill_text(line, Q), ts, False, TEXT) for line in it.get("lines", [])])
        for p in tb.text_frame.paragraphs:   # bold the "Label:" lead of each line
            t = p.runs[0].text
            if ":" in t:
                lab, rest = t.split(":", 1)
                p.runs[0].text = lab + ":"; p.runs[0].font.bold = True; p.runs[0].font.color.rgb = BLUE
                r = p.add_run(); r.text = rest; r.font.size = Pt(ts); r.font.color.rgb = TEXT; r.font.name = "Calibri"
            p.space_after = Pt(6 if n <= 3 else 4)


def layout_checklist(slide, spec, Q, icons):
    items = spec["items"]; n = len(items)
    gap = Inches(0.08); rh = (H - gap * (n - 1)) / n
    for k, it in enumerate(items):
        y = Y0 + k * (rh + gap); name, tint = STATUS[str(it["status"]).lower()]
        rounded_box(slide, X0, y, W, rh, fill=tint, radius=0.08)
        isz = min(Inches(0.4), rh - Inches(0.16))
        icon(slide, name, X0 + Inches(0.25), y + (rh - isz) / 2, isz, icons)
        textbox(slide, X0 + Inches(0.85), y, Inches(5.6), rh, [(fill_text(it["heading"], Q), 16, True, BLUE)], anchor=MSO_ANCHOR.MIDDLE)
        textbox(slide, X0 + Inches(6.5), y, W - Inches(6.6), rh, [(fill_text(it.get("caption", ""), Q), 15, False, TEXT)], anchor=MSO_ANCHOR.MIDDLE)


def layout_icon_rows(slide, spec, Q, icons, deck_dir):
    items = spec["items"]; n = len(items)
    img = spec.get("image"); iw = Inches(img.get("width", 4.2)) if img else 0
    tw = W if not img else W - iw - Inches(0.35)
    gap = Inches(0.2); rh = (H - gap * (n - 1)) / n
    hs = spec.get("heading_size", 20 if n <= 4 else 17); ts = spec.get("text_size", 14 if n <= 4 else 12.5)
    for k, it in enumerate(items):
        y = Y0 + k * (rh + gap); isz = min(Inches(0.9), rh - Inches(0.2))
        icon(slide, icon_file(it), X0 + Inches(0.1), y + (rh - isz) / 2, isz, icons)
        textbox(slide, X0 + Inches(1.25), y, tw - Inches(1.3), rh,
                [(fill_text(it["heading"], Q), hs, True, BLUE), (fill_text(it.get("caption", ""), Q), ts, False, TEXT)], anchor=MSO_ANCHOR.MIDDLE)
    if img:
        f = os.path.join(deck_dir, img["file"])
        pic = slide.shapes.add_picture(f, X0 + W - iw, Y0 + Inches(0.1), width=iw)
        if pic.height > H - Inches(0.2):
            ratio = (H - Inches(0.2)) / pic.height
            pic.width = int(pic.width * ratio); pic.height = int(H - Inches(0.2))


def layout_two_panel(slide, spec, Q, icons):
    panels = spec["panels"]; gap = Inches(0.5); pw = (W - gap) / 2
    for j, pn in enumerate(panels):
        x = X0 + j * (pw + gap); y = Y0
        rounded_box(slide, x, y, pw, H)
        hb = rounded_box(slide, x, y, pw, Inches(0.7), fill=BLUE, radius=0.12)
        if "icon" in pn:
            icon(slide, icon_file(pn, "white"), x + Inches(0.2), y + Inches(0.13), Inches(0.44), icons)
        textbox(slide, x + Inches(0.75), y, pw - Inches(0.85), Inches(0.7), [(fill_text(pn["title"], Q), 18, True, WHITE)], anchor=MSO_ANCHOR.MIDDLE)
        items = pn["items"]; n = len(items); rh = (H - Inches(0.9)) / n
        for k, it in enumerate(items):
            yy = y + Inches(0.8) + k * rh; isz = min(Inches(0.55), rh - Inches(0.15))
            tx = x + Inches(0.25)
            if "icon" in it:
                icon(slide, icon_file(it), x + Inches(0.25), yy + (rh - isz) / 2, isz, icons); tx = x + Inches(1.0)
            hs = spec.get("heading_size", 15); ts = spec.get("text_size", 12)
            paras = [(fill_text(it["heading"], Q), hs, True, BLUE)]
            if it.get("caption"):
                paras.append((fill_text(it["caption"], Q), ts, False, TEXT))
            textbox(slide, tx, yy, pw - (tx - x) - Inches(0.15), rh, paras, anchor=MSO_ANCHOR.MIDDLE)


def layout_grid(slide, spec, Q, icons):
    groups = spec["groups"]; ncol = spec.get("columns", 2); nrow = -(-len(groups) // ncol)
    HH = H - Inches(0.55) if spec.get("footnote") else H
    gx = Inches(0.35); gy = Inches(0.15); gw = (W - gx * (ncol - 1)) / ncol; gh = (HH - gy * (nrow - 1)) / nrow
    for k, g in enumerate(groups):
        r, c = divmod(k, ncol); x = X0 + c * (gw + gx); y = Y0 + r * (gh + gy)
        rounded_box(slide, x, y, gw, gh, radius=0.1)
        isz = min(Inches(0.6), gh - Inches(0.2))
        icon(slide, icon_file(g), x + Inches(0.15), y + (gh - isz) / 2, isz, icons)
        head = fill_text(g["title"], Q)
        if "count_of" in g:
            head += f"  ({sum(Q.get('domain_counts', {}).get(d, 0) for d in g['count_of'])})"
        textbox(slide, x + Inches(0.9), y, gw - Inches(1.0), gh,
                [(head, 15, True, BLUE), (fill_text(g.get("caption", ""), Q), 11.5, False, TEXT)], anchor=MSO_ANCHOR.MIDDLE)
    if spec.get("footnote"):   # small text under the grid (the content area is shortened to make room)
        textbox(slide, X0, Y0 + HH + Inches(0.05), W, Inches(0.5), [(fill_text(spec["footnote"], Q), 8, False, MUTED)])


def layout_title(slide, spec, Q, icons, deck_dir):
    """Title slide: shorten the subtitle box, add a row of logos and an acknowledgement line."""
    for sh in slide.placeholders:
        if sh.placeholder_format.type == 4:   # SUBTITLE
            sh.height = Inches(1.3)
        if sh.placeholder_format.type in (1, 3) and sh.has_text_frame and len(sh.text_frame.text) > 70:   # a long title: fit it
            for para in sh.text_frame.paragraphs:
                for run in para.runs:
                    run.font.size = Pt(spec.get("title_size", 30))
    logos = spec.get("logos", []); y = Inches(5.45); h = Inches(0.55)
    widths = [Inches(l.get("width", 2.0)) for l in logos]; gap = Inches(0.9)
    x = X0 + (W - sum(widths) - gap * (len(logos) - 1)) / 2
    for l, w in zip(logos, widths):
        pic = slide.shapes.add_picture(os.path.join(deck_dir, l["file"]), x, y, width=w)
        if pic.height > h:
            ratio = h / pic.height; pic.width = int(pic.width * ratio); pic.height = int(h)
        pic.top = int(y + (h - pic.height) / 2)
        x += w + gap
    if spec.get("acknowledgement"):
        textbox(slide, X0, Inches(6.1), W, Inches(0.4), [(fill_text(spec["acknowledgement"], Q), 12, False, MUTED)], align=PP_ALIGN.CENTER)


LAYOUTS = {"pipeline": layout_pipeline, "cards": layout_cards, "checklist": layout_checklist, "two_panel": layout_two_panel, "grid": layout_grid}


# ---- deck ---------------------------------------------------------------------------------
def clear_slides(prs):
    sldIdLst = prs.slides._sldIdLst
    for sldId in list(sldIdLst):
        prs.part.drop_rel(sldId.rId); sldIdLst.remove(sldId)


def draw(slide, s, Q, icons, deck_dir):
    if s["layout"] == "icon_rows":
        layout_icon_rows(slide, s, Q, icons, deck_dir)
    elif s["layout"] == "title":
        layout_title(slide, s, Q, icons, deck_dir)
    else:
        LAYOUTS[s["layout"]](slide, s, Q, icons)


def build_in_place(spec_path, pptx):
    """Draw the spec's slides on the rendered deck, matching each by title."""
    deck_dir = os.path.dirname(os.path.abspath(spec_path)); root = os.path.abspath(os.path.join(deck_dir, "..", ".."))
    spec = yaml.safe_load(open(spec_path, encoding="utf-8"))
    qmd = os.path.join(deck_dir, spec["deck"])
    Q = json.load(open(re.sub(r"[.]qmd$", ".quantities.json", qmd), encoding="utf-8"))
    icons = os.path.join(root, "docs", "slides", "img", "icons")
    prs = Presentation(pptx)
    by_title = {}
    for sl in prs.slides:
        if sl.shapes.title is not None and sl.slide_layout.name != "Section Header":
            by_title.setdefault(sl.shapes.title.text.strip(), sl)
    n = 0; missing = []
    for s in spec["slides"]:
        if s["layout"] == "title":
            slide = next((sl for sl in prs.slides if sl.slide_layout.name == "Title Slide"), None)
        else:
            slide = by_title.get(fill_text(s["title"], Q))
        if slide is None:
            missing.append(s["title"]); continue
        if s["layout"] != "title":
            for sh in list(slide.shapes):   # drop pandoc's empty body placeholder and anything else but the title
                if sh.is_placeholder and sh.placeholder_format.type in (1, 3, 13, 15, 16):   # title, centre title, slide number, footer, date
                    continue
                sh._element.getparent().remove(sh._element)
        draw(slide, s, Q, icons, deck_dir); n += 1
    nf = fill_figures(prs) if spec.get("fill_figures") else 0
    prs.save(pptx)
    print(f"{pptx}: {n} concept slide(s) drawn in place, {nf} lone figure(s) enlarged" + (f"; NOT FOUND: {missing}" if missing else ""))
    if missing:
        sys.exit(1)


def fill_figures(prs):
    """Enlarge the picture on every slide that holds one picture and no body text to the
    content area, keeping its aspect ratio (pandoc caps it at 4.75 in high)."""
    n = 0
    for sl in prs.slides:
        pics = [sh for sh in sl.shapes if sh.shape_type == 13]
        texts = [sh for sh in sl.shapes if sh.has_text_frame and sh.text_frame.text.strip()
                 and not (sh.is_placeholder and sh.placeholder_format.type in (1, 3, 13, 15, 16))]
        if len(pics) != 1 or texts or sl.slide_layout.name in ("Title Slide", "Section Header"):
            continue
        pic = pics[0]; r = min(W / pic.width, H / pic.height)
        pic.width = int(pic.width * r); pic.height = int(pic.height * r)
        pic.left = int(X0 + (W - pic.width) / 2); pic.top = int(Y0); n += 1
    return n


def build(spec_path):
    deck_dir = os.path.dirname(os.path.abspath(spec_path)); root = os.path.abspath(os.path.join(deck_dir, "..", ".."))
    spec = yaml.safe_load(open(spec_path, encoding="utf-8"))
    qmd = os.path.join(deck_dir, spec["deck"]); qmd_lines = open(qmd, encoding="utf-8").read().splitlines()
    Q = json.load(open(re.sub(r"[.]qmd$", ".quantities.json", qmd), encoding="utf-8"))
    ref = re.search(r'reference-doc:\s*"([^"]+)"', "\n".join(qmd_lines)).group(1)
    icons = os.path.join(root, "docs", "slides", "img", "icons")
    prs = Presentation(ref); clear_slides(prs)
    layout = next(l for l in prs.slide_layouts if l.name == "Title Only")
    rows = []
    for k, s in enumerate(spec["slides"], 1):
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = fill_text(s["title"], Q)
        draw(slide, s, Q, icons, deck_dir)
        notes = deck_notes(qmd_lines, s["title"], Q)
        slide.notes_slide.notes_text_frame.text = f"Replaces deck slide {s['replaces']} ({s['title']}).\n\n{notes}".strip()
        rows.append((k, s["replaces"], s["title"], s["layout"]))
    out = re.sub(r"[.]ya?ml$", ".pptx", spec_path); prs.save(out); print("wrote", out)
    write_placement(spec_path, rows, spec["deck"])
    previews(out, os.path.join(deck_dir, "concept_previews"))


def write_placement(spec_path, rows, deck):
    md = re.sub(r"[.]ya?ml$", "-PLACEMENT.md", spec_path)
    with open(md, "w", encoding="utf-8") as f:
        f.write(f"# Placement of the concept slides in {deck.replace('.qmd', '.pptx')}\n\n")
        f.write("Built by scripts/concept_slides/build.sh. Copy each slide from the concept deck over the\n"
                "deck slide it replaces (Home > New Slide > Reuse Slides, or copy-paste with *Use destination\n"
                "theme*); the speaker notes of the replaced slide are already on the new slide.\n\n")
        f.write("| Concept slide | Replaces deck slide | Title | Layout |\n|---|---|---|---|\n")
        for k, rep, title, lay in rows:
            f.write(f"| {k} | {rep} | {title} | {lay} |\n")
    print("wrote", md)


def previews(pptx, out_dir):
    os.makedirs(out_dir, exist_ok=True)
    soffice = r"C:\Program Files\LibreOffice\program\soffice.exe"
    if not os.path.exists(soffice):
        print("no LibreOffice: previews skipped"); return
    subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", out_dir, pptx], check=True, capture_output=True)
    pdf = os.path.join(out_dir, os.path.basename(pptx).replace(".pptx", ".pdf"))
    import fitz
    d = fitz.open(pdf)
    for f in os.listdir(out_dir):
        if f.endswith(".png"):
            os.remove(os.path.join(out_dir, f))
    for i, page in enumerate(d):
        page.get_pixmap(dpi=96).save(os.path.join(out_dir, f"slide{i + 1:02d}.png"))
    d.close(); os.remove(pdf)
    print(f"{len(os.listdir(out_dir))} preview(s) in {out_dir}")


if __name__ == "__main__":
    if "--into" in sys.argv:
        build_in_place(sys.argv[1], sys.argv[sys.argv.index("--into") + 1])
    else:
        build(sys.argv[1])
